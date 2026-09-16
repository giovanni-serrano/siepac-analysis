"""Referencia fija y recálculos: las pruebas nunca regeneran el valor esperado."""

from copy import deepcopy
import json

import numpy as np
import openpyxl
import pandas as pd
import pytest

from baseline_utils import CSV_RESULTADOS, LIBROS, comparar, leer_csv, leer_libro, leer_web
from config_siepac import ANIOS_ANALISIS as ANIOS, PAISES_SIEPAC as PAISES
from etl_soc2 import cargar_y_validar, calcular_regional
from generar_matriz_indicadores import calcular_valores, hoja_indicador, INDICADORES
from generar_visualizador import _fichas_visualizador
import generar_figuras_tesis as figuras
from procesar_dimensiones import leer_env, leer_soc, PESOS_SOC3
import procesar_dimensiones as dimensiones
from viz_comun import agregados_eco, construir_datos_json, preparar_datos


def base_eco_con_banderas(referencia):
    base = pd.DataFrame(referencia["csv"]["matriz_consolidada_wide.csv"])
    tarifas = pd.DataFrame(referencia["csv"]["tarifa_electrica_media.csv"])
    return base.merge(tarifas[["pais", "anio", "fuente_dato"]].rename(
        columns={"fuente_dato": "tarifa_fuente_dato"}),
        on=["pais", "anio"], validate="one_to_one")


@pytest.mark.parametrize("nombre", CSV_RESULTADOS)
def test_csv_actual_conserva_resultados(referencia, raiz_resultados, nombre):
    ruta = raiz_resultados / "data" / "processed" / nombre
    if not ruta.exists():
        pytest.skip("CSV local regenerable ausente: ejecutar pipeline aislado y usar --resultados-dir")
    comparar(leer_csv(ruta), referencia["csv"][nombre], nombre)


@pytest.mark.parametrize("nombre", LIBROS)
def test_libros_conservan_valores_formulas_y_notas(referencia, raiz_resultados, nombre):
    comparar(leer_libro(raiz_resultados / "data" / "processed" / nombre),
             referencia["libros"][nombre], nombre)


def test_web_conserva_todas_las_series_y_advertencias(referencia, raiz_resultados):
    comparar(leer_web(raiz_resultados / "graficos" / "visualizador_siepac.html"),
             referencia["web"], "web")


def test_recalculo_eco_desde_base_fija(referencia):
    base = base_eco_con_banderas(referencia)
    valores = calcular_valores(base)
    comparar(valores.to_dict("records"), referencia["csv"]["indicadores_ECO_valores.csv"], "ECO")


@pytest.mark.parametrize("codigo", ["ECO1", "ECO2", "ECO3", "ECO6", "ECO11", "ECO13", "ECO14", "ECO15"])
def test_agregados_eco_contra_referencia_fija(referencia, codigo):
    base = pd.DataFrame(referencia["csv"]["matriz_consolidada_wide.csv"])
    comparar(agregados_eco(base)[codigo], referencia["web"]["datos"][codigo]["agregado"], codigo)


def test_empaquetado_eco_conserva_promedios_paises_y_banderas(referencia):
    base = base_eco_con_banderas(referencia)
    valores = calcular_valores(base)
    hojas = {c: valores.pivot(index="pais", columns="anio", values=c).reset_index() for c in INDICADORES}
    hojas["tarifa_flag"] = valores[["pais", "anio", "tarifa_fuente_dato"]]
    actual = json.loads(construir_datos_json(preparar_datos(hojas), base, {}))
    for clave in [*INDICADORES, "imputados"]:
        comparar(actual[clave], referencia["web"]["datos"][clave], clave)


def test_soc2_recalcula_nacionales_componentes_y_regionales(referencia):
    df = cargar_y_validar()
    registros = df.astype(object).where(pd.notna(df), None).to_dict("records")
    comparar(registros, referencia["csv"]["soc2_pais_anio.csv"], "SOC2 nacional")
    comparar(calcular_regional(df).to_dict("records"), referencia["csv"]["soc2_regional.csv"], "SOC2 regional")


def test_cobertura_actual_30_claves_y_banderas_eco14(referencia, raiz_resultados):
    web = leer_web(raiz_resultados / "graficos" / "visualizador_siepac.html")
    assert web["anios"] == list(range(2020, 2025))
    assert set(web["paises"]) == set(PAISES)
    for clave, serie in web["datos"].items():
        if clave == "imputados":
            continue
        assert set(serie["paises"]) == set(PAISES), clave
        assert all(len(v) == 5 and all(np.isfinite(x) for x in v) for v in serie["paises"].values()), clave
    banderas = web["datos"]["imputados"]
    assert sum(sum(v) for v in banderas.values()) == 13
    assert all(v[-2:] == [True, True] for v in banderas.values())
    assert web["datos"]["ECO14"]["agregado"] is None
    assert "2023–2024" in web["fichas"]["ECO14"]["nota"]
    assert "imputad" in web["fichas"]["ECO14"]["nota"].lower()


@pytest.mark.parametrize("codigo", ["ECO1", "ECO2", "ECO3", "ECO6", "ECO11", "ECO13", "ECO14", "ECO15"])
def test_formulas_excel_generadas_conservan_referencia(referencia, codigo):
    libro = openpyxl.Workbook()
    base = referencia["csv"]["indicadores_ECO_valores.csv"]
    imputados = {(r["pais"], r["anio"]) for r in base if r["tarifa_fuente_dato"] == "imputado_CAGR"}
    hoja_indicador(libro, codigo, INDICADORES[codigo], imputados)
    filas = [list(f) for f in libro[codigo].iter_rows(values_only=True)]
    comparar(filas, referencia["libros"]["indicadores_ECO_SIEPAC.xlsx"][codigo], codigo)


def test_lectura_env_y_pib_ambiental_independiente(referencia, raiz_resultados):
    env = leer_env(raiz_resultados / "data" / "raw_equipo" / "ENVs.xlsx")
    publicado = referencia["web"]["datos"]
    for codigo in dimensiones.PESOS_AGREGADO:
        for pais in PAISES:
            comparar([round(float(v), 6) for v in env[codigo].loc[pais, ANIOS]],
                     publicado[codigo]["paises"][pais], f"{codigo}.{pais}")
    esperado = referencia["libros"]["indicadores_ENV_SIEPAC.xlsx"]["Datos_Base"]
    columnas = esperado[2]
    base_esperada = pd.DataFrame(esperado[3:], columns=columnas)
    np.testing.assert_allclose(env["BASE"].pib_usd_const2015, base_esperada.pib_usd_const2015, rtol=1e-12)
    eco = pd.DataFrame(referencia["csv"]["matriz_consolidada_wide.csv"])
    cruce = env["BASE"].merge(eco, on=["pais", "anio"], suffixes=("_env", "_eco"))
    assert not np.allclose(cruce.pib_usd_const2015_env, cruce.pib_usd_const2015_eco)


def test_lectura_soc_preserva_proxy_mix_uniforme(referencia, raiz_resultados, tmp_path, monkeypatch):
    # Usa el SOC2 congelado para que funcione también sin CSV locales ignorados.
    ruta = tmp_path / "soc2.csv"
    pd.DataFrame(referencia["csv"]["soc2_pais_anio.csv"]).to_csv(ruta, index=False)
    monkeypatch.setattr(dimensiones, "RUTA_SOC2", ruta)
    soc = leer_soc(raiz_resultados / "data" / "raw_equipo" / "SOCs.xlsx")
    for codigo in ("SOC1", "SOC2_PROM", "SOC2_VULNERABLE", *PESOS_SOC3):
        for pais in PAISES:
            comparar([round(float(v), 6) for v in soc[codigo].loc[pais, ANIOS]],
                     referencia["web"]["datos"][codigo]["paises"][pais], f"{codigo}.{pais}")
    for codigo, columna in (("SOC3_RURAL", "tasa_electrificacion_rural"),
                            ("SOC3_URB", "tasa_electrificacion_urbana")):
        for fila in soc["BASE"].to_dict("records"):
            esperado = fila[columna] * fila["pct_renovable_generacion"] / 100
            assert soc[codigo].loc[fila["pais"], fila["anio"]] == pytest.approx(esperado, abs=1e-9)


def test_fichas_actuales_conservan_unidades_formulas_y_limitaciones(referencia):
    comparar(_fichas_visualizador(), referencia["web"]["fichas"], "FICHAS")


def test_figuras_conservan_series_mediana_eco14_y_franja(referencia, tmp_path, monkeypatch):
    base = base_eco_con_banderas(referencia)
    valores = calcular_valores(base)
    hojas = {c: valores.pivot(index="pais", columns="anio", values=c).reset_index() for c in INDICADORES}
    hojas["tarifa_flag"] = valores[["pais", "anio", "tarifa_fuente_dato"]]
    hojas["datos_base"] = base
    datos = referencia["web"]["datos"]
    extra = {c: deepcopy(datos[c]) for c in figuras.SOC + figuras.ENV}
    extra["ENV6"] = {p: {"biomasa": datos["ENV6_BIOMASA"]["paises"][p],
                          "saldo": datos["ENV6_SALDO"]["paises"][p]} for p in PAISES}
    generadas = {}
    monkeypatch.setattr(figuras, "DIR_FIGURAS", tmp_path)
    monkeypatch.setattr(figuras, "cargar_datos", lambda: hojas)
    monkeypatch.setattr(figuras, "leer_series_extra", lambda: extra)
    monkeypatch.setattr(figuras, "_exportar", lambda fig, nombre: generadas.setdefault(nombre, fig))
    figuras.main()
    assert len(generadas) == 22
    for codigo in figuras.ECO + figuras.SOC + figuras.ENV:
        fig = generadas[f"{codigo}_bloque.png"]
        esperado = datos[codigo]["agregado"]
        if codigo == "ECO14":
            esperado = valores.pivot(index="anio", columns="pais", values=codigo).median(axis=1)
            assert "Mediana" in fig.data[3].name
            assert any(s.x0 == 2022.5 and s.x1 >= 2024 for s in fig.layout.shapes)
            assert any("CAGR" in a.text for a in fig.layout.annotations)
        np.testing.assert_allclose(fig.data[3].y, esperado, rtol=1e-12, atol=1e-9)
        np.testing.assert_allclose(fig.data[2].y, datos[codigo]["promedio"], rtol=1e-12, atol=5e-7)
    assert min(generadas["ECO15_bloque.png"].data[1].y) < 0
    env6 = generadas["ENV6_bloque.png"]
    np.testing.assert_allclose(env6.data[2].y, datos["ENV6_BIOMASA"]["agregado"], atol=1e-9)
    np.testing.assert_allclose(env6.data[5].y, datos["ENV6_SALDO"]["agregado"], atol=1e-9)


@pytest.mark.parametrize("cambio", ["valor", "signo", "bandera", "agregado", "cobertura"])
def test_comparador_detecta_mutaciones_cientificas(referencia, cambio):
    esperado = referencia["web"]["datos"]
    alterado = deepcopy(esperado)
    if cambio == "valor":
        alterado["ECO1"]["agregado"][0] += 1
    elif cambio == "signo":
        for pais, valores in alterado["ECO15"]["paises"].items():
            if any(v < 0 for v in valores):
                alterado["ECO15"]["paises"][pais] = [abs(v) for v in valores]
                break
    elif cambio == "bandera":
        alterado["imputados"][PAISES[0]][-1] = False
    elif cambio == "agregado":
        alterado["ECO14"]["agregado"] = alterado["ECO14"]["promedio"]
    else:
        alterado["SOC1"]["paises"][PAISES[0]].pop()
    with pytest.raises(AssertionError):
        comparar(alterado, esperado)
