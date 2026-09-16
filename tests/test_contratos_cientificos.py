"""Casos sintéticos independientes de las cifras de la tesis.

Los fallos conocidos expresan el rechazo deseado, sin normalizar el defecto
actual como comportamiento correcto ni modificar las funciones productivas.
"""

import json

import numpy as np
import openpyxl
import pandas as pd
import pytest

import consolidar_matriz as consolidacion
from config_siepac import ANIOS_ANALISIS as ANIOS, PAISES_SIEPAC as PAISES
import etl_consumo_final_total as consumo
import etl_consumo_industrial as industria
import etl_poblacion_total as poblacion
import etl_soc2 as soc2
import etl_tarifa_electrica_media as tarifa
from generar_matriz_indicadores import calcular_valores, hoja_indicador, INDICADORES
from generar_visualizador import _normalizar_env6
import procesar_dimensiones as dimensiones
from viz_comun import (agregados_eco, construir_datos_json, media_ponderada,
                       preparar_datos)


@pytest.fixture
def panel():
    return pd.DataFrame([{"pais": p, "anio": a, "valor_gwh": 1.0,
                          "valor_miles_hab": 1.0, "valor_habitantes": 1000.0}
                         for p in PAISES for a in ANIOS])


@pytest.fixture
def fuente_soc2():
    return pd.DataFrame([
        {"pais": p, "anio": a, "clientes_residenciales": 10 * (i + 1),
         "cargo_anual_usd": 100 + i * 30, "ingreso_prom_usd": 1000 + i * 1000,
         "ingreso_vulnerable_usd": 300 + i * 100,
         "proporcion_vulnerable": 0.296 if p == "Nicaragua" else 0.20,
         "fuente_hoja": "sintética", "nota_metodologica": "prueba",
         "observacion_clientes": "proxy"}
        for i, p in enumerate(PAISES) for a in ANIOS])


def cargar_soc2_temporal(df, tmp_path):
    ruta = tmp_path / "soc2.csv"
    df.to_csv(ruta, index=False)
    return soc2.cargar_y_validar(ruta)


@pytest.mark.parametrize("defecto", ["duplicado", "pais_ausente", "anio_ausente",
                                      "anio_fuera", "nulo", "infinito", "cero",
                                      "proporcion_nicaragua", "derivado_incorrecto"])
def test_soc2_rechaza_insumos_invalidos(fuente_soc2, tmp_path, defecto):
    df = fuente_soc2.copy()
    if defecto == "duplicado":
        # Mantiene 30 filas: prueba unicidad, no solo longitud.
        df.iloc[1] = df.iloc[0]
    elif defecto == "pais_ausente":
        df = df[df.pais != PAISES[0]]
    elif defecto == "anio_ausente":
        df = df[df.anio != ANIOS[-1]]
    elif defecto == "anio_fuera":
        df.loc[df.anio == ANIOS[-1], "anio"] = ANIOS[-1] + 1
    elif defecto in ("nulo", "infinito", "cero"):
        df["cargo_anual_usd"] = df["cargo_anual_usd"].astype(float)
        df.loc[0, "cargo_anual_usd"] = {"nulo": np.nan, "infinito": np.inf, "cero": 0}[defecto]
    elif defecto == "proporcion_nicaragua":
        df.loc[df.pais == "Nicaragua", "proporcion_vulnerable"] = 0.20
    else:
        df["soc2_prom_nacional"] = -999.0
    with pytest.raises(SystemExit, match="VALIDACIÓN FALLIDA"):
        cargar_soc2_temporal(df, tmp_path)


def test_soc2_razon_sumas_no_media_de_porcentajes(fuente_soc2, tmp_path):
    calculado = cargar_soc2_temporal(fuente_soc2, tmp_path)
    regional = soc2.calcular_regional(calculado)
    fila = fuente_soc2[fuente_soc2.anio == ANIOS[0]]
    clientes = fila.clientes_residenciales
    vulnerables = clientes * fila.proporcion_vulnerable
    esperado_prom = (fila.cargo_anual_usd * clientes).sum() / (fila.ingreso_prom_usd * clientes).sum() * 100
    esperado_vul = (fila.cargo_anual_usd * vulnerables).sum() / (fila.ingreso_vulnerable_usd * vulnerables).sum() * 100
    np.testing.assert_allclose(regional.SOC2_PROM, esperado_prom)
    np.testing.assert_allclose(regional.SOC2_VULNERABLE, esperado_vul)
    np.testing.assert_allclose(regional.brecha_vulnerable_prom_pp, esperado_vul - esperado_prom)
    porcentajes = fila.cargo_anual_usd / fila.ingreso_prom_usd * 100
    assert esperado_prom != pytest.approx(porcentajes.mean())
    assert esperado_prom != pytest.approx(np.average(porcentajes, weights=clientes))


@pytest.mark.parametrize("modulo", [consumo, industria, poblacion])
def test_etl_rechaza_pais_completamente_ausente(panel, modulo):
    with pytest.raises(SystemExit, match="VALIDACIÓN FALLIDA"):
        modulo.validar(panel[panel.pais != PAISES[0]])


@pytest.mark.parametrize("modulo", [consumo, industria, poblacion])
@pytest.mark.parametrize("defecto", ["duplicado", "anio_ausente"])
def test_etl_debe_rechazar_cobertura_invalida(panel, modulo, defecto):
    df = (pd.concat([panel, panel.iloc[[0]]], ignore_index=True)
          if defecto == "duplicado" else panel[panel.anio != ANIOS[-1]])
    with pytest.raises((SystemExit, ValueError), match="VALIDACIÓN FALLIDA"):
        modulo.validar(df)


@pytest.mark.parametrize("defecto", ["duplicado", "pais_ausente", "anio_ausente", "variable_ausente"])
def test_consolidacion_debe_rechazar_panel_incompleto(panel, defecto):
    datos = {}
    for nombre, columnas in consolidacion.MAPA_VARIABLES.items():
        df = panel[["pais", "anio"]].copy()
        df["fuente"] = "sintética"
        for columna in columnas:
            df[columna] = 1.0
        if defecto == "pais_ausente":
            df = df[df.pais != PAISES[0]]
        elif defecto == "anio_ausente":
            df = df[df.anio != ANIOS[-1]]
        datos[nombre] = df
    if defecto == "duplicado":
        df = datos["consumo_final_total.csv"]
        duplicada = df.iloc[[0]].copy()
        duplicada["valor_kwh"] = 999.0
        datos["consumo_final_total.csv"] = pd.concat([df, duplicada], ignore_index=True)
    elif defecto == "variable_ausente":
        del datos["consumo_final_total.csv"]
    with pytest.raises((SystemExit, ValueError), match="VALIDACIÓN FALLIDA"):
        tidy, wide = consolidacion.transformar(datos)
        consolidacion.validar(tidy, wide, datos)


@pytest.fixture
def base_eco():
    filas = []
    for i, p in enumerate(PAISES):
        for a in ANIOS:
            filas.append(dict(
                pais=p, anio=a, consumo_final_total_kwh=100 if i == 0 else 10,
                poblacion_habitantes=100 if i == 0 else 1,
                pib_usd_const2015=200, produccion_bruta_kwh=100,
                consumo_industrial_kwh=20, vai_usd_const2015=50,
                gen_fosil_kwh=40, gen_total_kwh=100,
                gen_hidro_kwh=10, gen_geotermia_kwh=10, gen_eolica_kwh=10,
                gen_solar_kwh=10, gen_biomasa_kwh=20,
                importaciones_kwh=0, exportaciones_kwh=20,
                tarifa_usd_mwh=150,
                tarifa_fuente_dato="imputado_CAGR" if a >= 2023 else "real"))
    return pd.DataFrame(filas)


def test_media_simple_y_agregado_eco_son_diferentes(base_eco):
    valores = calcular_valores(base_eco)
    valores["tarifa_imputada"] = valores.tarifa_fuente_dato == "imputado_CAGR"
    paquete = json.loads(construir_datos_json(valores, base_eco, {}))
    assert paquete["ECO1"]["promedio"] == [8.5] * len(ANIOS)
    np.testing.assert_allclose(paquete["ECO1"]["agregado"], [150 / 105] * len(ANIOS), atol=5e-7)
    # La aproximación ECO3 permanece limitada a consumo final / generación.
    np.testing.assert_allclose(valores.ECO3, base_eco.consumo_final_total_kwh)
    np.testing.assert_allclose(valores.ECO11 + valores.ECO13, 100)


def test_eco15_exportador_neto_conserva_negativos(base_eco):
    calculado = calcular_valores(base_eco)
    assert calculado.ECO15.tolist() == [-25.0] * len(base_eco)
    assert agregados_eco(base_eco)["ECO15"] == [-25.0] * len(ANIOS)


def test_eco14_sin_agregado_y_con_banderas(base_eco):
    valores = calcular_valores(base_eco)
    hojas = {codigo: valores.pivot(index="pais", columns="anio", values=codigo).reset_index()
             for codigo in INDICADORES}
    hojas["tarifa_flag"] = valores[["pais", "anio", "tarifa_fuente_dato"]]
    preparado = preparar_datos(hojas)
    paquete = json.loads(construir_datos_json(preparado, base_eco, {}))
    assert paquete["ECO14"]["agregado"] is None
    assert all(banderas == [False, False, False, True, True]
               for banderas in paquete["imputados"].values())
    libro = openpyxl.Workbook()
    imputados = {(p, a) for p in PAISES for a in ANIOS if a >= 2023}
    hoja_indicador(libro, "ECO14", INDICADORES["ECO14"], imputados)
    hoja = libro["ECO14"]
    assert not any(f[0].value == "Agregado regional (razón de sumas)" for f in hoja)
    assert "no definido" in hoja["A12"].value
    for fila in range(4, 10):
        assert hoja.cell(fila, 5).fill.fgColor.rgb == "00FFF2CC"
        assert hoja.cell(fila, 6).fill.fgColor.rgb == "00FFF2CC"


def test_cagr_usa_historia_y_no_modifica_observados():
    # El tramo anterior a 2020 determina un crecimiento anual exacto del 10 %.
    historia = pd.DataFrame([
        {"pais": p, "anio": a, "valor_usd_mwh": 100 * 1.1 ** (a - 2015)}
        for p in PAISES for a in [2015, 2020, 2021, 2022]])
    salida = tarifa.transformar(tarifa.proyectar_faltantes(historia))
    assert set(salida.anio) == set(ANIOS)
    assert len(salida) == 30
    np.testing.assert_allclose(salida.valor_usd_mwh, 100 * 1.1 ** (salida.anio - 2015))
    assert (salida.loc[salida.anio >= 2023, "fuente_dato"] == "imputado_CAGR").all()
    assert (salida.loc[salida.anio <= 2022, "fuente_dato"] == "real").all()
    observados = salida[salida.anio <= 2022].merge(historia, on=["pais", "anio"])
    assert (observados.valor_usd_mwh_x == observados.valor_usd_mwh_y).all()


def test_tarifa_debe_rechazar_anio_ausente(base_eco):
    df = base_eco.rename(columns={"tarifa_usd_mwh": "valor_usd_mwh",
                                 "tarifa_fuente_dato": "fuente_dato"})
    with pytest.raises((SystemExit, ValueError), match="VALIDACIÓN FALLIDA"):
        tarifa.validar(df[df.anio != ANIOS[-1]])


def test_conversion_energia_a_unidad_base(panel):
    panel["valor_gwh"] = 2.5
    salida = consumo.transformar(panel)
    assert (salida.valor_kwh == 2_500_000).all()


def test_env6_conserva_dos_magnitudes_y_signos():
    paquete = {"ENV6": {p: {"biomasa": [10.0] * len(ANIOS),
                                "saldo": [-3.0] * len(ANIOS)} for p in PAISES}}
    _normalizar_env6(paquete)
    assert set(paquete) == {"ENV6_BIOMASA", "ENV6_SALDO"}
    assert paquete["ENV6_BIOMASA"]["agregado"] == [60.0] * len(ANIOS)
    assert paquete["ENV6_SALDO"]["agregado"] == [-18.0] * len(ANIOS)
    assert paquete["ENV6_SALDO"]["paises"][PAISES[0]] == [-3.0] * len(ANIOS)


def test_ponderacion_por_denominador():
    valores = {p: [10.0 if i == 0 else 20.0] * len(ANIOS) for i, p in enumerate(PAISES)}
    pesos = {p: [9.0 if i == 0 else 1.0] * len(ANIOS) for i, p in enumerate(PAISES)}
    np.testing.assert_allclose(media_ponderada(valores, pesos), 190 / 14, atol=5e-7)


def test_dimensiones_usa_poblaciones_y_pib_propios(tmp_path, monkeypatch):
    """Ejercita main y sus ponderadores reales con libros temporales sintéticos."""
    serie = pd.DataFrame([[10.0 if i == 0 else 20.0] * len(ANIOS)
                          for i in range(len(PAISES))], index=PAISES, columns=ANIOS)
    base = pd.DataFrame([dict(pais=p, anio=a,
                             poblacion_miles=9 if i == 0 else 1,
                             pib_usd_const2015=1 if i == 0 else 9,
                             produccion_bruta_gwh=2 if i == 0 else 1)
                         for i, p in enumerate(PAISES) for a in ANIOS])
    env = {c: serie.copy() for c in dimensiones.PESOS_AGREGADO}
    env["BASE"] = base
    env["ENV6"] = pd.DataFrame([[p, s] + [1.0] * len(ANIOS)
                                 for p in PAISES for s in ("Inyección Biomasa", "Saldo MER")],
                                columns=["pais", "serie"] + ANIOS)
    soc = {c: serie.copy() for c in ("SOC1", "SOC2_PROM", "SOC2_VULNERABLE", "SOC3_RURAL", "SOC3_URB")}
    soc["BASE"] = base[["pais", "anio"]].copy()
    soc["BASE_SOC2"] = base[["pais", "anio"]].copy()
    zona = base[["pais", "anio"]].copy()
    zona["poblacion_rural_hab"] = base.poblacion_miles
    zona["poblacion_urbana_hab"] = base.pib_usd_const2015
    zona["poblacion_total_hab"] = zona.poblacion_rural_hab + zona.poblacion_urbana_hab
    zona.to_csv(tmp_path / "zona.csv", index=False)
    total = base[["pais", "anio"]].copy()
    total["valor_habitantes"] = [19 if p == PAISES[0] else 3 for p in base.pais]
    total.to_csv(tmp_path / "poblacion_total.csv", index=False)
    pd.DataFrame({"anio": ANIOS, "SOC2_PROM": [2.0] * len(ANIOS),
                  "SOC2_VULNERABLE": [12.0] * len(ANIOS)}).to_csv(tmp_path / "soc2.csv", index=False)
    monkeypatch.setattr(dimensiones, "leer_env", lambda ruta: env)
    monkeypatch.setattr(dimensiones, "leer_soc", lambda ruta: soc)
    for nombre, valor in {"DIR_OUT": tmp_path, "DIR_PROCESSED": tmp_path,
                          "RUTA_POB_ZONA": tmp_path / "zona.csv",
                          "RUTA_SOC2_REGIONAL": tmp_path / "soc2.csv"}.items():
        monkeypatch.setattr(dimensiones, nombre, valor)
    dimensiones.main()
    for dim, esperados in {
        "SOC": {"SOC1": 490 / 34, "SOC3_RURAL": 190 / 14, "SOC3_URB": 910 / 46,
                "SOC2_PROM": 2.0, "SOC2_VULNERABLE": 12.0},
        "ENV": {"ENV1_PC": 190 / 14, "ENV1_PIB": 910 / 46, "ENV3": 120 / 7},
    }.items():
        libro = openpyxl.load_workbook(tmp_path / f"indicadores_{dim}_SIEPAC.xlsx", data_only=True)
        try:
            for codigo, esperado in esperados.items():
                assert libro[codigo]["B10"].value == pytest.approx(110 / 6)
                assert libro[codigo]["B11"].value == pytest.approx(esperado, abs=5e-7)
        finally:
            libro.close()


def test_eco_debe_rechazar_denominador_cero(base_eco):
    base_eco.loc[0, "poblacion_habitantes"] = 0
    with pytest.raises((SystemExit, ValueError), match="VALIDACIÓN FALLIDA"):
        calcular_valores(base_eco)


@pytest.fixture
def libro_env_sintetico(tmp_path):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, ancho in {"ENV1": 7, "ENV2": 10, "ENV3": 11, "ENV6": 7}.items():
        libro.create_sheet(nombre).append([f"columna_{i}" for i in range(ancho)])
    libro["ENV1"]["E1"] = "C: PIB Real (Millones USD Constantes)"
    for p in PAISES:
        for a in ANIOS:
            # La intensidad no tiene valor cacheado; debe usar el PIB ambiental.
            libro["ENV1"].append([p, a, 4.0, 1000, 2, 0.004, None])
            libro["ENV2"].append([p, a, 1, 2, 1000, 2, 1, 2, 3, 4])
            libro["ENV3"].append([p, a, 100, 1, 2, 3, 4, 5, 6, 7, 8])
        libro["ENV6"].append([p, "Inyección Biomasa"] + [10] * len(ANIOS))
        libro["ENV6"].append([None, "Saldo MER"] + [-3] * len(ANIOS))
    ruta = tmp_path / "env.xlsx"
    libro.save(ruta)
    libro.close()
    return ruta


def test_env_preserva_pib_del_libro_sin_sustituirlo(libro_env_sintetico):
    salida = dimensiones.leer_env(libro_env_sintetico)
    assert (salida["BASE"].pib_usd_const2015 == 2).all()


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason="FUERA DE ALCANCE (D06): respaldo ENV1 sin caché interpreta PIB en millones como USD")
def test_env1_sin_cache_debe_conservar_escala(libro_env_sintetico):
    salida = dimensiones.leer_env(libro_env_sintetico)
    # 4 miles de t = 4 millones de kg; 2 millones de USD => 2 kg/USD.
    np.testing.assert_allclose(salida["ENV1_PIB"], 2.0, rtol=1e-12)


def test_lectura_env_debe_rechazar_duplicados(libro_env_sintetico):
    libro = openpyxl.load_workbook(libro_env_sintetico)
    libro["ENV1"].append([PAISES[0], ANIOS[0], 999, 1000, 2, 0.999, None])
    libro.save(libro_env_sintetico)
    libro.close()
    with pytest.raises((SystemExit, ValueError), match="VALIDACIÓN FALLIDA"):
        dimensiones.leer_env(libro_env_sintetico)
