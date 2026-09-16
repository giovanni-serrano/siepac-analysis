"""Fase 2: límites de validación y rechazo antes de escribir productos."""

import numpy as np
import openpyxl
import pandas as pd
import pytest

import consolidar_matriz as consolidacion
from config_siepac import ANIOS_ANALISIS as ANIOS, PAISES_SIEPAC as PAISES
from etl_comun import validar_panel
import etl_tarifa_electrica_media as tarifa
import generar_matriz_indicadores as eco
import procesar_dimensiones as dimensiones
from test_contratos_cientificos import base_eco, libro_env_sintetico, panel


def test_panel_detecta_hueco_aunque_estan_todos_los_paises_y_anios(panel):
    with pytest.raises(SystemExit, match=r"claves país-año faltantes.*Costa Rica.*2020"):
        validar_panel(panel.iloc[1:], "panel sintético")


@pytest.mark.parametrize("anio", [2020.5, np.inf, "sin año"])
def test_panel_rechaza_anios_no_enteros(panel, anio):
    panel["anio"] = panel["anio"].astype(object)
    panel.loc[0, "anio"] = anio
    with pytest.raises(SystemExit, match=r"años inválidos.*Costa Rica"):
        validar_panel(panel, "panel sintético")


def test_panel_detecta_duplicado_con_anio_textual_equivalente(panel):
    otra = panel.iloc[[0]].copy()
    otra["anio"] = otra["anio"].astype(str)
    with pytest.raises(SystemExit, match=r"duplicadas.*Costa Rica.*2020"):
        validar_panel(pd.concat([panel, otra], ignore_index=True), "panel sintético")


def test_panel_admite_anios_textuales_sin_mutar_entrada(panel):
    panel["anio"] = panel["anio"].astype(str)
    antes = panel.copy(deep=True)
    validar_panel(panel, "panel sintético")
    pd.testing.assert_frame_equal(panel, antes)


@pytest.mark.parametrize("hoja", ["ENV1", "ENV2", "ENV3"])
@pytest.mark.parametrize("defecto", ["duplicado", "hueco", "pais_ausente", "anio_ausente"])
def test_env_rechaza_claves_invalidas_en_cada_hoja(libro_env_sintetico, hoja, defecto):
    libro = openpyxl.load_workbook(libro_env_sintetico)
    ws = libro[hoja]
    if defecto == "duplicado":
        ws.append([c.value for c in ws[2]])
    elif defecto == "hueco":
        ws.delete_rows(2)
    else:
        for fila in range(ws.max_row, 1, -1):
            if ((defecto == "pais_ausente" and ws.cell(fila, 1).value == PAISES[0])
                    or (defecto == "anio_ausente" and ws.cell(fila, 2).value == ANIOS[-1])):
                ws.delete_rows(fila)
    libro.save(libro_env_sintetico)
    libro.close()
    with pytest.raises(SystemExit, match=rf"VALIDACIÓN FALLIDA.*{hoja}"):
        dimensiones.leer_env(libro_env_sintetico)


@pytest.mark.parametrize("columna", ["poblacion_habitantes", "pib_usd_const2015",
                                     "produccion_bruta_kwh", "vai_usd_const2015", "gen_total_kwh"])
@pytest.mark.parametrize("valor", [0, -1, np.nan, np.inf, -np.inf])
def test_eco_rechaza_denominadores_invalidos_con_clave(base_eco, columna, valor):
    base_eco[columna] = base_eco[columna].astype(float)
    base_eco.loc[0, columna] = valor
    with pytest.raises(SystemExit, match=rf"VALIDACIÓN FALLIDA.*{columna}.*Costa Rica.*2020"):
        eco.calcular_valores(base_eco)


@pytest.mark.parametrize("exportaciones", [100, 101])
def test_eco15_rechaza_disponibilidad_no_positiva(base_eco, exportaciones):
    base_eco.loc[0, "exportaciones_kwh"] = exportaciones
    with pytest.raises(SystemExit, match=r"ECO15.*disponibilidad_kwh.*Costa Rica"):
        eco.calcular_valores(base_eco)


def test_eco_invalido_no_sobrescribe_ninguna_salida(base_eco, tmp_path, monkeypatch):
    base_eco.loc[0, "poblacion_habitantes"] = 0
    wide = tmp_path / "wide.csv"
    tidy = tmp_path / "tidy.csv"
    base_eco.drop(columns="tarifa_fuente_dato").to_csv(wide, index=False)
    base_eco[["pais", "anio", "tarifa_fuente_dato"]].rename(
        columns={"tarifa_fuente_dato": "fuente_dato"}).assign(
            variable="tarifa_usd_mwh").to_csv(tidy, index=False)
    salida_excel = tmp_path / "indicadores.xlsx"
    salida_csv = tmp_path / "indicadores.csv"
    for ruta in (salida_excel, salida_csv):
        ruta.write_bytes(b"salida anterior intacta")
    for nombre, ruta in {"RUTA_WIDE": wide, "RUTA_TIDY": tidy,
                         "RUTA_SALIDA": salida_excel, "RUTA_VALORES": salida_csv}.items():
        monkeypatch.setattr(eco, nombre, ruta)
    with pytest.raises(SystemExit, match="VALIDACIÓN FALLIDA"):
        eco.main()
    assert salida_excel.read_bytes() == b"salida anterior intacta"
    assert salida_csv.read_bytes() == b"salida anterior intacta"


def test_consolidacion_archivo_ausente_no_genera_salidas(tmp_path, monkeypatch):
    monkeypatch.setattr(consolidacion, "DIR_PROCESSED", tmp_path)
    with pytest.raises(SystemExit, match=r"archivo requerido ausente.*consumo_final_total.csv"):
        consolidacion.main()
    assert list(tmp_path.iterdir()) == []


def test_tarifa_historia_duplicada_falla_antes_de_cagr(base_eco):
    historia = base_eco[["pais", "anio", "tarifa_usd_mwh"]].rename(
        columns={"tarifa_usd_mwh": "valor_usd_mwh"})
    historia = pd.concat([historia, historia.iloc[[0]]], ignore_index=True)
    with pytest.raises(SystemExit, match=r"tarifa histórica.*duplicadas.*Costa Rica.*2020"):
        tarifa.proyectar_faltantes(historia)


def test_tarifa_historia_insuficiente_no_publica_panel_incompleto():
    historia = pd.DataFrame([dict(pais=p, anio=2020, valor_usd_mwh=100.0) for p in PAISES])
    parcial = tarifa.transformar(tarifa.proyectar_faltantes(historia))
    with pytest.raises(SystemExit, match=r"tarifa final / ECO14.*años faltantes"):
        tarifa.validar(parcial)


@pytest.mark.parametrize("bandera", [None, "observado", ""])
def test_tarifa_rechaza_banderas_ausentes_o_desconocidas(base_eco, bandera):
    df = base_eco.rename(columns={"tarifa_usd_mwh": "valor_usd_mwh",
                                 "tarifa_fuente_dato": "fuente_dato"})
    df.loc[0, "fuente_dato"] = bandera
    with pytest.raises(SystemExit, match=r"fuente_dato inválida.*Costa Rica.*2020"):
        tarifa.validar(df)
