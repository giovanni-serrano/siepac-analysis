"""
datos_dimensiones.py — Lectura y preparación científica ENV/SOC
====================================================
Etapa del pipeline : resultados estructurados (módulo importable)
Entradas           : ENVs.xlsx, SOCs.xlsx y CSV normalizados de población/SOC2
Salidas            : DataFrames de bases, series y agregados en memoria
Alimenta           : libros, resumen, tablas, figuras y visualizador
Fuente de datos    : ETL normalizados y libros originales del equipo
Uso                : importar funciones; invocadas por las etapas existentes
Notas metodológicas: mantiene fuentes, ponderadores, orden de operaciones y
                     precisión publicada. No sustituye insumos originales.
Autor: Luis Giovanni Serrano Bello — Tesis SIEPAC, UNI Nicaragua
"""

import logging
import sys
from pathlib import Path
import openpyxl
import pandas as pd
from config_siepac import PAISES_SIEPAC as PAISES, ANIOS_ANALISIS as ANIOS
from etl_comun import fallar_validacion, validar_panel
from calculos_indicadores import media_ponderada
from metadatos_indicadores import PESOS_AGREGADO, PESOS_SOC3
log = logging.getLogger(__name__)

def leer_env(ruta: Path) -> dict[str, pd.DataFrame]:
    """Devuelve {clave_serie: DataFrame pais x anio} + ENV6 especial."""
    wb = openpyxl.load_workbook(ruta, data_only=True)   # valores, no formulas
    # Revisar las claves originales antes de int(), pivot o diccionarios:
    # ninguna fila repetida puede sustituir a otra ni ocultar un hueco.
    try:
        for nombre in ("ENV1", "ENV2", "ENV3"):
            if nombre not in wb.sheetnames:
                fallar_validacion(f"ENV / {ruta.name}", f"hoja faltante: {nombre}")
            claves = pd.DataFrame(
                [(f[0], f[1]) for f in wb[nombre].iter_rows(min_row=2, values_only=True)],
                columns=["pais", "anio"])
            validar_panel(claves, f"{ruta.name} / {nombre}")
    except BaseException:
        wb.close()
        raise
    series = {}

    def tabla(hoja, col_valor, recalcular=None):
        ws = wb[hoja]
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        reg = {}
        for f in filas:
            pais, anio = f[0], int(f[1])
            v = f[col_valor]
            if v is None and recalcular:
                v = recalcular(f)
            reg[(pais, anio)] = float(v)
        df = pd.DataFrame(
            [[reg[(p, a)] for a in ANIOS] for p in PAISES],
            index=PAISES, columns=ANIOS)
        return df

    # ENV1: col F (idx 5) per cápita; col G (idx 6) intensidad almacenada.
    series["ENV1_PC"] = tabla("ENV1", 5)
    # Miles de t -> kg y millones de USD -> USD aportan ambos 10⁶:
    # (f[2] × 10⁶) / (f[4] × 10⁶) = f[2] / f[4], en kg CO2eq/USD.
    series["ENV1_PIB"] = tabla("ENV1", 6,
        recalcular=lambda f: f[2] / f[4])
    # ENV2: 4 salidas (cols G..J -> idx 6..9)
    series["ENV2_SO2_PC"] = tabla("ENV2", 6)
    series["ENV2_PAR_PC"] = tabla("ENV2", 7)
    series["ENV2_SO2_PIB"] = tabla("ENV2", 8)
    series["ENV2_PAR_PIB"] = tabla("ENV2", 9)
    # ENV3: escala electrica g/kWh (ultima col, idx 10)
    series["ENV3"] = tabla("ENV3", 10)

    # ENV6: pares de filas (Inyeccion Biomasa / Saldo MER) por pais
    ws = wb["ENV6"]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    env6 = []
    pais_actual = None
    for f in filas:
        if f[0]:
            pais_actual = str(f[0]).strip()
        serie = str(f[1]).strip()
        env6.append([pais_actual, serie] + [float(v) for v in f[2:7]])
    series["ENV6"] = pd.DataFrame(
        env6, columns=["pais", "serie"] + ANIOS)

    # Datos base de la dimension: variables de entrada de ENV1/ENV2/ENV3.
    reg = {}
    for f in wb["ENV1"].iter_rows(min_row=2, values_only=True):
        reg[(f[0], int(f[1]))] = {
            "emisiones_gei_10e3t": float(f[2]),
            "poblacion_miles": float(f[3]),
            "pib_usd_const2015": float(f[4])}
    for f in wb["ENV2"].iter_rows(min_row=2, values_only=True):
        reg[(f[0], int(f[1]))].update({
            "so2_10e3t": float(f[2]), "particulas_10e3t": float(f[3])})
    for f in wb["ENV3"].iter_rows(min_row=2, values_only=True):
        reg[(f[0], int(f[1]))].update({
            "produccion_bruta_gwh": float(f[2]),
            "nox_10e3t": float(f[4]), "co_10e3t": float(f[5])})
    base = pd.DataFrame(
        [{"pais": p, "anio": a, **reg[(p, a)]}
         for p in PAISES for a in ANIOS])
    series["BASE"] = base
    wb.close()
    return series


def leer_soc(ruta: Path, ruta_soc2: Path) -> dict[str, pd.DataFrame]:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    series = {}

    # --- SOC1: anios en filas (descendentes), paises en columnas C..H ---
    ws = wb["SOC 1"]
    orden_cols = ["Costa Rica", "El Salvador", "Guatemala",
                  "Honduras", "Nicaragua", "Panamá"]  # fila 2, cols C..H
    reg = {}
    for f in ws.iter_rows(min_row=3, max_row=7, values_only=True):
        anio = int(f[0])
        for j, pais in enumerate(orden_cols):
            reg[(pais, anio)] = float(f[2 + j])
    series["SOC1"] = pd.DataFrame(
        [[reg[(p, a)] for a in ANIOS] for p in PAISES],
        index=PAISES, columns=ANIOS)

    # --- SOC2: salida normalizada del ETL específico ------------------
    if not ruta_soc2.exists():
        sys.exit("VALIDACIÓN FALLIDA — falta soc2_pais_anio.csv; ejecutar "
                 "antes etl_soc2.py.")
    soc2 = pd.read_csv(ruta_soc2)
    esperadas = set(PAISES) | {""}
    if len(soc2) != 30 or soc2.duplicated(["pais", "anio"]).any():
        sys.exit("VALIDACIÓN FALLIDA — SOC2 procesado no contiene 30 "
                 "claves país-año únicas.")
    if set(soc2["pais"]) != esperadas - {""}:
        sys.exit("VALIDACIÓN FALLIDA — países inesperados en SOC2.")
    for clave in ("SOC2_PROM", "SOC2_VULNERABLE"):
        series[clave] = (soc2.pivot(index="pais", columns="anio",
                                    values=clave)
                         .reindex(index=PAISES, columns=ANIOS))

    # --- SOC3: tabla tidy oculta en columnas K..Y ---
    # Se filtra por contenido (col K = pais valido, col L = anio) porque
    # la hoja mezcla bloques auxiliares y el encabezado cambia de fila.
    ws = wb["SOC 3"]
    reg_r, reg_u = {}, {}
    n_filas = 0
    for f in ws.iter_rows(values_only=True):
        if f[10] in PAISES and f[11] is not None:
            try:
                anio = int(f[11])
            except (TypeError, ValueError):
                continue
            if anio not in ANIOS:
                continue
            pais = f[10]
            reg_r[(pais, anio)] = float(f[23])  # col X: acceso renov. rural
            reg_u[(pais, anio)] = float(f[24])  # col Y: acceso renov. urbano
            n_filas += 1
    assert n_filas == 30, f"SOC3: se esperaban 30 filas pais-anio, hay {n_filas}"

    # Datos base de la dimensión: tasas de electrificación y % renovable
    # (SOC1/SOC3), más los insumos unitarios en USD y ponderadores de SOC2.
    reg_b = {}
    for f in wb["SOC 3"].iter_rows(values_only=True):
        if f[10] in PAISES and f[11] is not None:
            try:
                anio = int(f[11])
            except (TypeError, ValueError):
                continue
            if anio not in ANIOS:
                continue
            reg_b[(f[10], anio)] = {
                "pct_renovable_generacion": float(f[20]),
                "tasa_electrificacion_rural": float(f[21]),
                "tasa_electrificacion_urbana": float(f[22])}
    for clave, reg in [("SOC3_RURAL", reg_r), ("SOC3_URB", reg_u)]:
        series[clave] = pd.DataFrame(
            [[reg[(p, a)] for a in ANIOS] for p in PAISES],
            index=PAISES, columns=ANIOS)

    series["BASE"] = pd.DataFrame(
        [{"pais": p, "anio": a,
          "pct_sin_electricidad": float(series["SOC1"].loc[p, a]),
          **reg_b[(p, a)]}
         for p in PAISES for a in ANIOS])
    columnas_soc2 = [
        "pais", "anio", "clientes_residenciales", "cargo_anual_usd",
        "ingreso_prom_usd", "ingreso_vulnerable_usd",
        "proporcion_vulnerable", "clientes_vulnerables_proxy",
        "gasto_prom_proxy_usd", "ingreso_prom_proxy_usd",
        "gasto_vulnerable_proxy_usd", "ingreso_vulnerable_proxy_usd",
        "fuente_dato",
    ]
    series["BASE_SOC2"] = soc2[columnas_soc2].copy()
    return series


def _valores_de(df: pd.DataFrame) -> dict:
    """DataFrame pais x anio -> {pais: [v_2020..v_2024]} con None en NaN."""
    return {p: [None if pd.isna(df.loc[p, a]) else float(df.loc[p, a])
                for a in ANIOS] for p in PAISES}


def _pesos_de(base: pd.DataFrame, col: str) -> dict:
    """Columna de la tabla BASE (tidy pais-anio) -> {pais: [w_2020..]}."""
    return {p: [float(base.loc[(base["pais"] == p) & (base["anio"] == a),
                               col].iloc[0]) for a in ANIOS]
            for p in PAISES}


def preparar_dimensiones(env, soc, dir_processed, ruta_pob_zona, ruta_soc2_regional):
    """Completa las bases y construye ambos agregados sin escribir productos."""
    # Poblacion total (peso del agregado de SOC1), producida por el ETL.
    ruta_pob = dir_processed / "poblacion_total.csv"
    pesos_pob = None
    if ruta_pob.exists():
        pob = pd.read_csv(ruta_pob)
        pesos_pob = _pesos_de(pob.rename(
            columns={"valor_habitantes": "poblacion"}), "poblacion")
    else:
        log.warning("Sin %s: SOC1 quedará sin fila de agregado regional.",
                    ruta_pob.name)

    if not ruta_pob_zona.exists():
        sys.exit("VALIDACIÓN FALLIDA — falta poblacion_rural_urbana.csv; "
                 "ejecutar antes etl_poblacion_rural_urbana.py.")
    pob_zona = pd.read_csv(ruta_pob_zona)
    pesos_soc3 = {
        clave: _pesos_de(pob_zona, columna)
        for clave, columna in PESOS_SOC3.items()
    }
    columnas_zona = ["pais", "anio", "poblacion_rural_hab",
                     "poblacion_urbana_hab", "poblacion_total_hab"]
    soc["BASE"] = soc["BASE"].merge(
        pob_zona[columnas_zona], on=["pais", "anio"], how="left",
        validate="one_to_one")
    if soc["BASE"][columnas_zona[2:]].isna().any().any():
        sys.exit("VALIDACIÓN FALLIDA — faltan ponderadores SOC3 tras "
                 "cruzar población rural/urbana.")
    soc["BASE"] = soc["BASE"].merge(
        soc["BASE_SOC2"], on=["pais", "anio"], how="left",
        validate="one_to_one")

    if not ruta_soc2_regional.exists():
        sys.exit("VALIDACIÓN FALLIDA — falta soc2_regional.csv; ejecutar "
                 "antes etl_soc2.py.")
    soc2_regional = pd.read_csv(ruta_soc2_regional).set_index("anio")
    agregados_soc2 = {
        clave: [float(soc2_regional.loc[a, clave]) for a in ANIOS]
        for clave in ("SOC2_PROM", "SOC2_VULNERABLE")
    }

    agregados_env = {
        clave: media_ponderada(_valores_de(env[clave]),
                              _pesos_de(env["BASE"], columna))
        for clave, columna in PESOS_AGREGADO.items()
    }
    agregados_soc = {}
    for clave in ("SOC1", "SOC2_PROM", "SOC2_VULNERABLE", *PESOS_SOC3):
        if clave == "SOC1" and pesos_pob:
            agregado = media_ponderada(_valores_de(soc[clave]), pesos_pob)
        elif clave in pesos_soc3:
            agregado = media_ponderada(_valores_de(soc[clave]), pesos_soc3[clave])
        elif clave in agregados_soc2:
            agregado = agregados_soc2[clave]
        else:
            agregado = None
        agregados_soc[clave] = agregado
    return agregados_env, agregados_soc
